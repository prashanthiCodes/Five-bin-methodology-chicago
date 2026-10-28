from __future__ import annotations
import csv, json, math, shutil, statistics
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
YEARS=(2014,2019,2024)
VINTAGE={2014:2010,2019:2010,2024:2020}
LABELS=("Very Low","Low","Middle","High","Very High")
NEUTRAL=LABELS
PRICE=ROOT/"data/source/Tract_Level_Sales_Weighted_Audit.xlsx"

SPECS=[
 dict(id="Bachelor_or_higher_attainment_share",label="Bachelor's Degree or Higher Attainment Share",kind="share",direction="positive",n="Pop_Edu_BachAbove_E",d="Total_Pop_25Over_E",unit="% of population aged 25+",formula="Population aged 25+ with a bachelor's degree or higher ÷ total population aged 25+ × 100"),
 dict(id="Median_household_income",label="Median Household Income",kind="median",direction="positive",src="HH_Med_Inc_E",weight="Total_HH_E.x",unit="USD per household (source dollar basis)",formula="Median(household income)"),
 dict(id="Median_house_price",label="Median House Price",kind="price",direction="positive",boundary_vintage=2020,unit="USD assumed",formula="Σ (housing-type median price × sales count) ÷ Σ sales counts"),
 dict(id="Owner_occupied_share",label="Owner-Occupied Housing Share",kind="share",direction="positive",n="Total_Own_HH_E.x",d="Total_HH_E.x",unit="% of occupied housing units",formula="Owner-occupied units ÷ total occupied units × 100"),
 dict(id="Income_100k_plus_share",label="Households with Income $100,000 or More",kind="share",direction="positive",n="HH_Inc_100K_More_E",d="Total_HH_E.x",unit="% of households",formula="Households with income ≥ $100,000 ÷ total households × 100"),
 dict(id="HU_Built_2000_Later_E",label="Housing Units Built 2000 or Later",kind="count",direction="positive",src="HU_Built_2000_Later_E",unit="housing units",formula="Count of housing units built 2000 or later"),
 dict(id="Age_65_over_share",label="Population Aged 65 or Older",kind="share",direction="neutral",n="Pop_Age_65Over_E",d="Total_Pop_E",unit="% of total population",formula="Population aged 65+ ÷ total population × 100"),
 dict(id="Age_10_17_share",label="Population Aged 10–17",kind="share",direction="neutral",n="Pop_Age_10to17_E",d="Total_Pop_E",unit="% of total population",formula="Population aged 10–17 ÷ total population × 100"),
 dict(id="Below_Poverty_Total_Pop_Share",label="Population Below Poverty",kind="share",direction="negative",n="Pop_Below_Pov_E",d="Total_Pop_E",unit="% of total population",formula="People below poverty ÷ total tract population × 100"),
 dict(id="HU_Built_1939_Earlier",label="Housing Units Built 1939 or Earlier",kind="count",direction="negative",src="HU_Built_1939_Earlier_E",unit="housing units",formula="Count of housing units built 1939 or earlier"),
 dict(id="Total_Vac_HU_E",label="Vacant Housing Units",kind="count",direction="negative",src="Total_Vac_HU_E",unit="housing units",formula="Count of vacant housing units"),
 dict(id="HH_Cost_Burden_Share",label="Cost-Burdened Households",kind="share",direction="negative",nf=["Own_HH_Cost_Burden_E","Rent_HH_Cost_Burden_E"],df=["Total_Own_HH_IncComp_E","Total_Rent_HH_IncComp_E.x"],unit="% of households with applicable cost data",formula="(Owner + renter cost-burdened households) ÷ households with applicable cost data × 100"),
 dict(id="Renter_Occupied_Share",label="Renter-Occupied Housing Share",kind="share",direction="neutral",n="Total_Rent_HH_E.x",d="Total_HH_E.x",unit="% of occupied housing units",formula="Renter-occupied units ÷ total occupied units × 100"),
]
PUBLIC=[s for s in SPECS if s["kind"]!="price"]

def num(row,key):
 if not key or key not in row or row[key] is None or str(row[key]).strip()=="": return None
 try: x=float(row[key])
 except (ValueError,TypeError): return None
 return x if math.isfinite(x) else None
def read_csv(path):
 with path.open(encoding="utf-8-sig",newline="") as f: return list(csv.DictReader(f))
def write_csv(path,rows,fields=None):
 path.parent.mkdir(parents=True,exist_ok=True)
 if fields is None: fields=list(rows[0]) if rows else []
 with path.open("w",encoding="utf-8-sig",newline="") as f:
  w=csv.DictWriter(f,fieldnames=fields,extrasaction="ignore"); w.writeheader(); w.writerows(rows)
def wmedian(pairs):
 p=sorted((v,w) for v,w in pairs if w and w>0)
 if not p:return None
 t=sum(w for _,w in p)/2; c=0
 for v,w in p:
  c+=w
  if c>=t:return v
 return p[-1][0]

def prices():
 if not PRICE.exists(): raise FileNotFoundError(f"Missing required workbook: {PRICE}")
 book=load_workbook(PRICE,read_only=True,data_only=True); allp={}; geos={}; audits=[]
 required={"GEOID","COMMUNITY_NM","Neighbourhood Median House Value","Missing_From_Sales","SFH_Sales_Count","TwoFour_Sales_Count","SFH_Median_Hp","TwoFour_Median_Hp"}
 try:
  for y in YEARS:
   ws=book[str(y)]; h=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
   missing=required-set(h)
   if missing: raise ValueError(f"{y}: missing columns {sorted(missing)}")
   ix={x:h.index(x) for x in required}; rec={}; dup=[]; bad=[]; miss=0; checked=0; mismatch=0
   ws.reset_dimensions()
   for row in ws.iter_rows(min_row=2,max_col=len(h),values_only=True):
    if not row or all(v is None for v in row): continue
    raw=str(row[ix["GEOID"]]).strip(); g=raw[:-2] if raw.endswith(".0") else raw
    if len(g)!=11 or not g.isdigit() or not g.startswith("17031"): bad.append(raw); continue
    if g in rec: dup.append(g); continue
    val=num({"v":row[ix["Neighbourhood Median House Value"]]},"v")
    flag=str(row[ix["Missing_From_Sales"]]).strip().lower()=="true"
    sf=num({"v":row[ix["SFH_Sales_Count"]]},"v") or 0
    tf=num({"v":row[ix["TwoFour_Sales_Count"]]},"v") or 0
    sm=num({"v":row[ix["SFH_Median_Hp"]]},"v")
    tm=num({"v":row[ix["TwoFour_Median_Hp"]]},"v")
    available=[(v,w) for v,w in ((sm,sf),(tm,tf)) if v is not None and w>0]
    if val and val>0 and available:
     checked+=1; expected=sum(v*w for v,w in available)/sum(w for _,w in available)
     if not math.isclose(val,expected,rel_tol=1e-9,abs_tol=.01): mismatch+=1
    if flag or val is None or val<=0: val=None; miss+=1
    rec[g]={"value":val,"support":sf+tf,"community":str(row[ix["COMMUNITY_NM"]] or "")}
   if dup or bad: raise ValueError(f"{y} price GEOID issues: duplicates={len(dup)} malformed={len(bad)}")
   vals=[r["value"] for r in rec.values() if r["value"] is not None]
   med=statistics.median(vals) if vals else None
   city=f"CITYWIDE MEDIAN {y}"; cached=None
   if city in h:
    first=next(ws.iter_rows(min_row=2,max_row=2,max_col=len(h),values_only=True),None)
    try: cached=float(first[h.index(city)]) if first else None
    except (TypeError,ValueError): pass
   allp[y]=rec; geos[y]=set(rec)
   audits.append(dict(year=y,worksheet=str(y),records=len(rec),unique_geoids=len(set(rec)),duplicate_geoids=len(dup),malformed_geoids=len(bad),missing_or_flagged=miss,valid_price_count=len(vals),minimum=min(vals) if vals else None,maximum=max(vals) if vals else None,unweighted_mean=statistics.fmean(vals) if vals else None,computed_median=med,workbook_cached_city_median=cached,cached_matches_computed=cached is not None and med is not None and math.isclose(cached,med,rel_tol=1e-9,abs_tol=.01),weight_formula_checked=checked,weight_formula_mismatches=mismatch,currency="Unspecified; USD assumed",metric_note="Sales-count-weighted combination of SFH and 2–4-unit medians; not pooled transaction median."))
 finally: book.close()
 return allp,audits,geos
def measure(s,row,price):
 if s["kind"]=="price":
  v=None if price is None else price["value"]
  return v,None,None,None if price is None else price["support"],"missing" if v is None else "valid"
 if s["kind"]=="median":
  v=num(row,s["src"]); wt=num(row,s["weight"])
  if v is None:return None,None,None,wt,"missing"
  return (None,None,None,wt,"invalid") if v<0 else (v,None,None,wt,"valid")
 if s["kind"]=="count":
  v=num(row,s["src"])
  if v is None:return None,None,None,None,"missing"
  return (None,None,None,None,"invalid") if v<0 else (v,None,None,None,"valid")
 if "nf" in s:
  ns=[num(row,k) for k in s["nf"]]; ds=[num(row,k) for k in s["df"]]
  if any(v is None for v in ns+ds):return None,None,None,None,"missing"
  n,d=sum(ns),sum(ds)
 else:
  n,d=num(row,s.get("n")),num(row,s.get("d"))
  if n is None or d is None:return None,n,d,None,"missing"
 if d<=0 or n<0 or n>d:return None,n,d,None,"invalid"
 return n/d,n,d,None,"valid"

def shown(v,kind): return None if v is None else v*100 if kind=="share" else v
def classify(rel,direction):
 raw=1 if rel<50 else 2 if rel<75 else 3 if rel<=125 else 4 if rel<=150 else 5
 i=raw-1
 if direction=="negative":i=4-i
 return raw,(NEUTRAL if direction=="neutral" else LABELS)[i]
def benchmark(s,items):
 if not items:return None,"No valid values",0,None,None
 if s["kind"]=="share":
  n=sum(x["n"] for x in items); d=sum(x["d"] for x in items)
  return (n/d if d else None),f"Ratio of summed numerator to summed denominator across valid Chicago tracts ({s.get('n',s.get('nf'))} / {s.get('d',s.get('df'))})",len(items),n,d
 if s["kind"]=="count":return statistics.fmean(x["v"] for x in items),"Unweighted arithmetic mean of valid tract counts",len(items),None,None
 if s["kind"]=="median":
  pairs=[(x["v"],x["wt"]) for x in items if x["wt"] and x["wt"]>0]
  return wmedian(pairs),"Household-weighted median of tract medians; not a direct Chicago-place ACS median",len(pairs),None,None
 if s["kind"]=="price":
  return statistics.median(x["v"] for x in items),"Median of valid tract workbook measures; not transaction-level Chicago median",len(items),None,None
 return None,"No benchmark",0,None,None

def main():
 inputs=ROOT/"data/source"; tables=ROOT/"outputs/tables"; private=ROOT/"outputs/internal"
 tables.mkdir(parents=True,exist_ok=True); private.mkdir(parents=True,exist_ok=True)
 pdat,paudit,pgeo=prices()
 sources={}; geos={}; duplicates={}; malformed={}; bounds={}
 for y in YEARS:
  rows=read_csv(inputs/f"acs_tract_inputs_{y}.csv"); sources[y]=rows
  seen=set(); duplicates[y]=[]; malformed[y]=[]
  for r in rows:
   g=str(r.get("geoid","")).strip()
   if len(g)!=11 or not g.isdigit() or not g.startswith("17031"):malformed[y].append(g)
   if g in seen:duplicates[y].append(g)
   seen.add(g)
  geos[y]={g for g in seen if len(g)==11 and g.isdigit() and g.startswith("17031")}
  v=VINTAGE[y]; path=ROOT/f"data/boundaries/chicago_tracts_{v}.geojson"
  features=json.loads(path.read_text(encoding="utf-8"))["features"]; ids=[f["properties"]["GEOID"] for f in features]
  if len(ids)!=len(set(ids)):raise ValueError(f"Duplicate boundary GEOIDs {path}")
  bounds[v]=set(ids)
 write_csv(private/"house_price_workbook_audit.csv",paudit)
 join_audit=[]
 for y in YEARS:
  unmatched=pgeo[y]-geos[y]; missing=geos[y]-pgeo[y]; geo_unmatched=pgeo[y]-bounds[2020]
  join_audit.append(dict(year=y,workbook_geoids=len(pgeo[y]),source_geoids=len(geos[y]),price_geoids_not_in_2020_boundary=len(geo_unmatched),boundary_2020_geoids_without_price=len(bounds[2020]-pgeo[y]),price_geoids_not_in_same_year_acs=len(unmatched),same_year_acs_geoids_without_price=len(missing),unmatched_workbook_ids="|".join(sorted(unmatched)),source_ids_without_price_row="|".join(sorted(missing)),price_ids_without_2020_boundary="|".join(sorted(geo_unmatched)),missing_price_values=sum(x["value"] is None for x in pdat[y].values()),geometry_assessment="2020 vintage is inferred: workbook IDs match 2020 geometry; workbook does not declare vintage."))
 write_csv(private/"house_price_geoid_join_audit.csv",join_audit)
 p_union=set().union(*(pgeo[y] for y in YEARS)); p_shared=set.intersection(*(pgeo[y] for y in YEARS))
 pmembership=[]
 for g in sorted(p_union):
  ys=[y for y in YEARS if g in pgeo[y]]
  pmembership.append(dict(GEOID=g,in_2014=2014 in ys,in_2019=2019 in ys,in_2024=2024 in ys,years_present="|".join(map(str,ys)),years_count=len(ys),shared_by_all_three=len(ys)==3))
 write_csv(private/"house_price_geoid_membership.csv",pmembership)
 summaries=[]; classifications=[]; quality=[]; alljson={}
 for y in YEARS:
  rows=sources[y]; v=VINTAGE[y]
  quality += [
   dict(year=y,check="Source GEOID validation",status="PASS" if not duplicates[y] and not malformed[y] else "WARN",count=len(geos[y]),details=f"Rows {len(rows)}; duplicates {len(duplicates[y])}; malformed {len(malformed[y])}."),
   dict(year=y,check="Boundary join",status="PASS" if geos[y]==bounds[v] else "WARN",count=len(geos[y]&bounds[v]),details=f"{v}-vintage polygons; data without boundary {len(geos[y]-bounds[v])}; boundary without data {len(bounds[v]-geos[y])}."),
   dict(year=y,check="Geographic vintage",status="INFO",count=len(bounds[v]),details=f"{y} uses {v} TIGER vintage. Same GEOID does not prove same polygon; no crosswalk/harmonization."),
   dict(year=y,check="House-price blanks",status="INFO",count=sum(x["value"] is None for x in pdat[y].values()),details="Flagged/zero/blank prices remain blank; no values copied between years."),
   dict(year=y,check="House-price geometry join",status="PASS" if pgeo[y] <= bounds[2020] else "WARN",count=len(pgeo[y] & bounds[2020]),details=f"House-price GEOIDs are joined to 2020-vintage geometry for all workbook years. 2020 boundary IDs without workbook rows: {len(bounds[2020]-pgeo[y])}. Workbook vintage is inferred from the ID match, not stated in the workbook."),
   dict(year=y,check="Renter count enrichment",status="WARN" if any(num(r,"Total_Rent_HH_E.x") is None for r in rows) else "PASS",count=sum(num(r,"Total_Rent_HH_E.x") is not None for r in rows),details="Extracted from the local master workbook after all shared tract fields matched exactly. Owner + renter equals total occupied housing units on matched rows; missing renter counts remain blank.")
  ]
  j={r["geoid"]:{"_meta":{"tract_name":r.get("tract_name",""),"community_area":r.get("community_area","")}} for r in rows}
  for s in SPECS:
   v_ind=s.get("boundary_vintage",v)
   rows_for=rows
   if s["kind"]=="price":
    rows_for=[{"geoid":g,"tract_name":"Census tract "+g[-6:],"community_area":pdat[y][g]["community"]} for g in sorted(pdat[y])]
    for pr in rows_for:
     j.setdefault(pr["geoid"],{"_meta":{"tract_name":pr["tract_name"],"community_area":pr["community_area"]}})
   vals={}; items=[]; nmiss=ninvalid=0
   for r in rows_for:
    g=r["geoid"]; price=pdat[y].get(g) if s["kind"]=="price" else None
    value,n,d,aux,status=measure(s,r,price); vals[g]=(value,n,d,aux,status,r)
    if status=="missing":nmiss+=1
    elif status=="invalid":ninvalid+=1
    else:items.append(dict(v=value,n=n,d=d,wt=aux))
   b,definition,nb,bn,bd=benchmark(s,items)
   if b is not None and (not math.isfinite(b) or b<=0):b=None
   rawvals=[x["v"] for x in items]
   summaries.append(dict(year=y,indicator=s["id"],label=s["label"],direction=s["direction"],unit=s["unit"],
    minimum_tract_value=shown(min(rawvals) if rawvals else None,s["kind"]),
    maximum_tract_value=shown(max(rawvals) if rawvals else None,s["kind"]),
    unweighted_arithmetic_mean_tract_value=shown(statistics.fmean(rawvals) if rawvals else None,s["kind"]),
    median_tract_value=shown(statistics.median(rawvals) if rawvals else None,s["kind"]),
    benchmark_value=shown(b,s["kind"]),benchmark_definition=definition,benchmark_numerator=bn,benchmark_denominator=bd,
    valid_tract_count=len(items),missing_tract_count=nmiss,invalid_tract_count=ninvalid,total_source_tract_count=len(rows_for),tracts_used_in_benchmark=nb))
   quality.append(dict(year=y,check=f"{s['id']} valid/missing/invalid",status="WARN" if nmiss or ninvalid else "PASS",count=len(items),details=f"missing={nmiss}; invalid={ninvalid}."))
   if s["id"]=="Bachelor_or_higher_attainment_share":
    quality.append(dict(year=y,check="Bachelor's-or-higher source mapping",status="REVIEW",count=len(items),details="Uses supplied Pop_Edu_BachAbove_E / Total_Pop_25Over_E fields. Includes postgraduate degrees. The source workbook does not provide ACS table IDs/codebook; 2024 tracts with zero 25+ denominator remain invalid."))
   if s["id"]=="HH_Cost_Burden_Share":
    quality.append(dict(year=y,check="Cost-burden source mapping",status="REVIEW",count=len(items),details="Uses owner+renter burden count / owner+renter income-complete households. Source workbook lacks table IDs/codebook for full verification."))
   if s["id"]=="Renter_Occupied_Share":
    quality.append(dict(year=y,check="Renter-share source mapping",status="REVIEW",count=len(items),details="Uses renter units / total occupied units. Source workbook lacks table IDs/codebook; mapping conforms to requested B25003 definition."))
   for g,(value,n,d,aux,srcstatus,r) in vals.items():
    rel=value/b*100 if srcstatus=="valid" and b and b>0 else None
    raw,label=classify(rel,s["direction"]) if rel is not None else (None,None)
    status="classified" if rel is not None else "missing_value" if srcstatus=="missing" else "invalid_value" if srcstatus=="invalid" else "missing_benchmark"
    note=("Bachelor's degree or higher; includes master's, professional, and doctoral degrees. Source field table IDs are not documented in the supplied workbook." if s["id"]=="Bachelor_or_higher_attainment_share" else
          "Sales-count-weighted average of SFH and 2–4-unit medians; not a pooled median; USD assumed." if s["kind"]=="price" else
          "Descriptive measure; higher or lower is not inherently favorable or unfavorable." if s["direction"]=="neutral" else "")
    rec=dict(year=y,geoid=g,tract_name=r.get("tract_name",""),community_area=r.get("community_area",""),boundary_vintage=v_ind,
     indicator=s["id"],indicator_label=s["label"],direction=s["direction"],unit=s["unit"],tract_value=shown(value,s["kind"]),
     numerator_count=n,denominator_count=d,supporting_count=aux,benchmark_value=shown(b,s["kind"]),relative_percent=rel,
     raw_bin=raw,class_label=label,status=status,benchmark_definition=definition,quality_note=note)
    classifications.append(rec)
    j[g][s["id"]]=dict(value=rec["tract_value"],numerator_count=n,denominator_count=d,supporting_count=aux,benchmark=rec["benchmark_value"],relative_percent=rel,raw_bin=raw,class_label=label,status=status,quality_note=note)
  alljson[y]=j
  write_csv(tables/f"classification_{y}.csv",[r for r in classifications if r["year"]==y])
 union=set().union(*(geos[y] for y in YEARS)); shared=set.intersection(*(geos[y] for y in YEARS)); membership=[]
 for g in sorted(union):
  ys=[y for y in YEARS if g in geos[y]]
  membership.append(dict(GEOID=g,in_2014=2014 in ys,in_2019=2019 in ys,in_2024=2024 in ys,years_present="|".join(map(str,ys)),years_count=len(ys),shared_by_all_three=len(ys)==3,only_one_year=len(ys)==1,two_years=len(ys)==2))
 write_csv(private/"geoid_membership_all_years.csv",membership)
 exc=[]
 for y in YEARS:
  exc += [dict(year=y,GEOID=g,issue="duplicate_source_GEOID") for g in duplicates[y]]
  exc += [dict(year=y,GEOID=g,issue="malformed_source_GEOID") for g in malformed[y]]
  exc += [dict(year=y,GEOID=g,issue="data_without_matching_boundary") for g in sorted(geos[y]-bounds[VINTAGE[y]])]
  exc += [dict(year=y,GEOID=g,issue="boundary_without_data") for g in sorted(bounds[VINTAGE[y]]-geos[y])]
 write_csv(private/"geoid_validation_exceptions.csv",exc,["year","GEOID","issue"])
 gs=[]
 for y in YEARS:
  v=VINTAGE[y]
  gs.append(dict(year=y,source_rows=len(sources[y]),unique_valid_geoids=len(geos[y]),duplicates=len(duplicates[y]),malformed=len(malformed[y]),shared_by_all_three=len(shared),
   only_one_year_in_this_year=sum(x["only_one_year"] and x[f"in_{y}"] for x in membership),two_years_including_this_year=sum(x["two_years"] and x[f"in_{y}"] for x in membership),
   boundary_vintage=v,boundary_geoids=len(bounds[v]),data_without_boundary=len(geos[y]-bounds[v]),boundary_without_data=len(bounds[v]-geos[y]),
   note="2014/2019 use 2010 and 2024 uses 2020 vintage. Same GEOID does not establish identical polygon; no crosswalk or harmonization."))
 write_csv(private/"geoid_validation_summary.csv",gs)
 write_csv(tables/"classification_all_years.csv",classifications)
 write_csv(tables/"summary_statistics_by_indicator_year.csv",summaries)
 write_csv(tables/"indicator_value_extremes_by_year.csv",[dict(year=r["year"],indicator=r["indicator"],label=r["label"],unit=r["unit"],least_valid_tract_value=r["minimum_tract_value"],highest_valid_tract_value=r["maximum_tract_value"],valid_tract_count=r["valid_tract_count"],missing_tract_count=r["missing_tract_count"],invalid_tract_count=r["invalid_tract_count"],total_source_tract_count=r["total_source_tract_count"]) for r in summaries])
 write_csv(tables/"city_benchmarks.csv",[dict(year=r["year"],indicator=r["indicator"],label=r["label"],direction=r["direction"],benchmark_value=r["benchmark_value"],unit=r["unit"],definition=r["benchmark_definition"],valid_tracts=r["valid_tract_count"],tracts_used_in_benchmark=r["tracts_used_in_benchmark"]) for r in summaries])
 write_csv(tables/"data_quality_report.csv",quality)
 meta=dict(indicators=SPECS,bins=[dict(raw_bin=i+1,range=["Below 50%","50% to below 75%","75% through 125% inclusive","Above 125% through 150% inclusive","Above 150%"][i],positive_label=LABELS[i],negative_label=LABELS[4-i],neutral_label=NEUTRAL[i]) for i in range(5)],
  benchmark_note="Relative percent = tract value / same-year benchmark × 100. Share benchmarks are pooled numerator/denominator rates; count benchmarks are unweighted mean tract counts; median benchmarks are documented tract-distribution references.",
  geography_note="2014/2019 use 2010-vintage tracts; 2024 uses 2020-vintage tracts. Not a fixed-geography panel.")
 internal=ROOT/"website/assets/data"; public=ROOT/"docs/assets/data"
 internal.mkdir(parents=True,exist_ok=True); public.mkdir(parents=True,exist_ok=True)
 (internal/"indicator_metadata.json").write_text(json.dumps(meta,indent=2,ensure_ascii=False),encoding="utf-8")
 publicmeta=dict(meta); publicmeta["indicators"]=PUBLIC; publicmeta["publication_note"]="Review build. Price workbook values withheld pending sharing approval."
 (public/"indicator_metadata.json").write_text(json.dumps(publicmeta,indent=2,ensure_ascii=False),encoding="utf-8")
 for y in YEARS:
  (internal/f"classification_{y}.json").write_text(json.dumps(alljson[y],separators=(",",":"),ensure_ascii=False),encoding="utf-8")
  public_vintage=2020 if y==2024 else 2010
  public_geo=json.loads((ROOT/f"data/boundaries/chicago_tracts_{public_vintage}.geojson").read_text(encoding="utf-8"))
  allowed={feature["properties"]["GEOID"] for feature in public_geo["features"]}
  pub={g:{k:({fk:fv for fk,fv in v.items() if fk!="quality_note"} if k!="_meta" else v) for k,v in r.items() if k!="Median_house_price"} for g,r in alljson[y].items() if g in allowed}
  (public/f"classification_{y}.json").write_text(json.dumps(pub,separators=(",",":"),ensure_ascii=False),encoding="utf-8")
  for v in (2010,2020):
   src=ROOT/f"data/boundaries/chicago_tracts_{v}.geojson"; dst=ROOT/f"docs/data/boundaries/chicago_tracts_{v}.geojson"
   dst.parent.mkdir(parents=True,exist_ok=True)
   if not dst.exists():shutil.copy2(src,dst)
 print(json.dumps(dict(tracts={str(y):len(geos[y]) for y in YEARS},shared_all_three=len(shared),indicators=len(SPECS),public_indicators=len(PUBLIC),summary_rows=len(summaries),price_audit=paudit),indent=2,ensure_ascii=False))
if __name__=="__main__":main()
