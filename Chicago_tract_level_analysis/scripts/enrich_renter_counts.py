import csv, json, math, shutil
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]
SOURCES={2014:("acs_5yr_2014_master.xlsx","Selected tracts"),2019:("acs_5yr_2019_master.xlsx","Sheet1"),2024:("acs_5yr_2024_master.xlsx","Sheet1")}
META={"year","geoid","tract_name","community_area","community_area_id","source_geoid_column","source_workbook"}

def clean(v):
 if v is None:return ""
 t=str(v).strip()
 return t[:-2] if t.endswith(".0") else t

def same(a,b):
 if a in (None,"") and b in (None,""):return True
 try:return math.isclose(float(a),float(b),rel_tol=0,abs_tol=0)
 except (ValueError,TypeError):return str(a).strip()==str(b).strip()

def run():
 audit=[]
 source_dir=ROOT/"data/source"; workbook_dir=ROOT/"data/raw/acs_master_inputs"; legacy=source_dir/"legacy"
 legacy.mkdir(parents=True,exist_ok=True)
 for year,(filename,sheet_name) in SOURCES.items():
  path=source_dir/f"acs_tract_inputs_{year}.csv"
  rows=list(csv.DictReader(path.open(encoding="utf-8-sig",newline="")))
  if not rows:raise ValueError(f"No source rows for {year}")
  headers=list(rows[0])
  if "Total_Rent_HH_E.x" not in headers:headers.append("Total_Rent_HH_E.x")
  wb=load_workbook(workbook_dir/filename,read_only=True,data_only=True)
  try:
   if sheet_name not in wb.sheetnames:raise ValueError(f"{year} source sheet missing: {sheet_name}")
   ws=wb[sheet_name]; original_headers=[c.value for c in next(ws.iter_rows(min_row=1,max_row=1))]
   needed={"GEOID","Total_Rent_HH_E.x","Total_Own_HH_E.x","Total_HH_E.x"}
   if not needed.issubset(set(original_headers)):raise ValueError(f"{year} missing renter mapping fields: {needed-set(original_headers)}")
   ix={h:original_headers.index(h) for h in original_headers if h is not None}; ws.reset_dimensions()
   lookup={}; duplicates=[]
   for raw in ws.iter_rows(min_row=2,max_col=len(original_headers),values_only=True):
    if not raw or all(v is None for v in raw):continue
    geoid=clean(raw[ix["GEOID"]])
    if not geoid:continue
    if geoid in lookup:duplicates.append(geoid)
    lookup[geoid]=raw
   if duplicates:raise ValueError(f"{year} duplicate source GEOIDs: {duplicates[:5]}")
   if not (legacy/f"acs_tract_inputs_{year}_before_renter.csv").exists():
    shutil.copy2(path,legacy/f"acs_tract_inputs_{year}_before_renter.csv")
   mismatches=[]; matched=0; missing=[]
   common=[field for field in rows[0] if field not in META and field in ix]
   for row in rows:
    g=row["geoid"]; raw=lookup.get(g)
    if raw is None:
     row["Total_Rent_HH_E.x"]=""; missing.append(g); continue
    matched+=1
    for field in common:
     if not same(row.get(field),raw[ix[field]]):mismatches.append((g,field))
    rent=raw[ix["Total_Rent_HH_E.x"]]; own=raw[ix["Total_Own_HH_E.x"]]; total=raw[ix["Total_HH_E.x"]]
    if rent is not None and own is not None and total is not None and not math.isclose(float(rent)+float(own),float(total),rel_tol=0,abs_tol=1):
     mismatches.append((g,"owner+renter!=total"))
    row["Total_Rent_HH_E.x"]="" if rent is None else rent
   if mismatches:raise ValueError(f"{year} source workbook mismatches: {mismatches[:8]}")
   with path.open("w",encoding="utf-8-sig",newline="") as f:
    writer=csv.DictWriter(f,fieldnames=headers);writer.writeheader();writer.writerows(rows)
   audit.append({"year":year,"source_rows":len(rows),"matched_rows":matched,"source_rows_missing_renter_value":len(missing),"source_geoids_missing_from_master":"|".join(missing),"master_geoids_not_in_source":len(set(lookup)-{r["geoid"] for r in rows}),"duplicate_master_geoids":len(duplicates),"preexisting_field_value_mismatches":len(mismatches),"matched_fields_checked":"|".join(common),"own_plus_renter_equals_total":"PASS on matched rows","mapping":"Total_Rent_HH_E.x / Total_HH_E.x"})
  finally:wb.close()
 out=ROOT/"outputs/internal/renter_source_join_audit.csv";out.parent.mkdir(parents=True,exist_ok=True)
 with out.open("w",encoding="utf-8-sig",newline="") as f:
  w=csv.DictWriter(f,fieldnames=list(audit[0]));w.writeheader();w.writerows(audit)
 print(json.dumps(audit,indent=2))
if __name__=="__main__":run()
